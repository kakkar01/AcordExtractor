"""Compare QA-corrected JSON vs run JSON from an Excel mapping.

Supported Excel row formats (any one works):
1) qa_json_path + run_json_path
2) qa_json + run_json (JSON text directly in cells)
3) file_name (or id) with --qa-dir and --run-dir (will resolve <name>.json)
4) blob_uri + final_extraction with --run-dir (resolves *_extracted.json)

Outputs:
- Console summary with total files, matched files, and total diff count.
- CSV with per-row diff_count and top differences.
- JSON with detailed per-row diffs.

Usage example:
    python scripts/compare_json_from_excel.py \
        --excel data/outputs/acord140_compare.xlsx \
        --sheet Sheet1 \
        --qa-col qa_json_path \
        --run-col run_json_path \
        --out-csv data/outputs/acord140_diff_report.csv \
        --out-json data/outputs/acord140_diff_report.json
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook


IGNORED_META_KEYS = {
    "_markdown",
    "_ocr_key_values",
    "_doc_intel_key_values",
    "_validation_warnings",
    "_field_reasons",
    "_raw",
    "_schema_mismatch",
}


@dataclass
class RowResult:
    row_number: int
    key: str
    qa_source: str
    run_source: str
    diff_count: int
    wrong_values: int
    right_values: int
    status: str
    top_differences: str


def _normalize_key(name: str) -> str:
    stem = Path(name).stem
    return stem.strip().lower()


def _to_json(value: Any) -> dict[str, Any] | list[Any] | str | int | float | bool | None:
    if isinstance(value, (dict, list, str, int, float, bool)) or value is None:
        return value
    return str(value)


def _read_json_from_path(path_text: str) -> Any:
    path = Path(path_text)
    if not path.exists():
        raise FileNotFoundError(f"JSON file not found: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _read_json_from_cell(text: str) -> Any:
    return json.loads(text)


def _path_join_json(base_dir: Path, key: str) -> Path:
    candidate = base_dir / key
    if candidate.suffix.lower() != ".json":
        candidate = candidate.with_suffix(".json")
    return candidate


def _blob_uri_to_output_candidates(run_dir: Path, blob_uri: str) -> list[Path]:
    """Generate likely run output JSON filenames from blob URI."""
    parsed = urlparse(blob_uri)
    path_parts = [unquote(p) for p in parsed.path.split("/") if p]
    if not path_parts:
        return []

    candidates: list[Path] = []

    # Match downloader naming: join last up-to-4 segments.
    tail = path_parts[-4:] if len(path_parts) >= 4 else path_parts
    joined_name = "_".join(tail)
    if joined_name:
        if not joined_name.lower().endswith(".pdf"):
            joined_name = joined_name + ".pdf"
        candidates.append(run_dir / f"{Path(joined_name).stem}_extracted.json")

    # Fallback: use only last file segment.
    last_name = path_parts[-1]
    if not last_name.lower().endswith(".pdf"):
        last_name = last_name + ".pdf"
    candidates.append(run_dir / f"{Path(last_name).stem}_extracted.json")

    return candidates


def _resolve_run_json_from_blob_uri(run_dir: Path, blob_uri: str) -> Path:
    candidates = _blob_uri_to_output_candidates(run_dir, blob_uri)
    for candidate in candidates:
        if candidate.exists():
            return candidate

    raise FileNotFoundError(
        "Could not find output JSON for blob_uri. "
        f"Checked candidates: {[str(p) for p in candidates]}"
    )


def _collect_diffs(qa: Any, run: Any, path: str = "$") -> list[str]:
    diffs: list[str] = []

    qa, run = _align_for_compare(qa, run)

    if _values_equivalent(qa, run):
        return diffs

    if type(qa) is not type(run):
        diffs.append(f"{path}: type mismatch qa={type(qa).__name__} run={type(run).__name__}")
        return diffs

    if isinstance(qa, dict):
        qa_map = {
            _normalize_field_key(k): k
            for k in qa.keys()
            if k not in IGNORED_META_KEYS
        }
        run_map = {
            _normalize_field_key(k): k
            for k in run.keys()
            if k not in IGNORED_META_KEYS
        }

        for missing_norm in sorted(set(qa_map.keys()) - set(run_map.keys())):
            missing_in_run = qa_map[missing_norm]
            diffs.append(f"{path}.{missing_in_run}: missing in run")

        # Schema-values-only comparison: ignore fields that exist only in run JSON.
        # This prevents metadata/output-only keys from being counted as differences.

        for shared_norm in sorted(set(qa_map.keys()) & set(run_map.keys())):
            qa_key = qa_map[shared_norm]
            run_key = run_map[shared_norm]
            child_path = f"{path}.{qa_key}"
            diffs.extend(_collect_diffs(qa[qa_key], run[run_key], child_path))
        return diffs

    if isinstance(qa, list):
        if len(qa) != len(run):
            diffs.append(f"{path}: list length mismatch qa={len(qa)} run={len(run)}")

        max_len = min(len(qa), len(run))
        for idx in range(max_len):
            child_path = f"{path}[{idx}]"
            diffs.extend(_collect_diffs(qa[idx], run[idx], child_path))
        return diffs

    if qa != run:
        diffs.append(f"{path}: value mismatch qa={qa!r} run={run!r}")

    return diffs


def _align_for_compare(qa: Any, run: Any) -> tuple[Any, Any]:
    """Normalize common wrapper mismatch: singleton list vs object."""
    if isinstance(qa, list) and len(qa) == 1 and isinstance(run, dict) and isinstance(qa[0], dict):
        return qa[0], run
    if isinstance(run, list) and len(run) == 1 and isinstance(qa, dict) and isinstance(run[0], dict):
        return qa, run[0]

    # ACORD 140 row-wise normalization: QA may store premises as list,
    # while run output stores object with `rows`.
    if isinstance(qa, list) and isinstance(run, dict) and isinstance(run.get("rows"), list):
        return qa, run.get("rows")
    if isinstance(run, list) and isinstance(qa, dict) and isinstance(qa.get("rows"), list):
        return qa.get("rows"), run

    return qa, run


def _count_value_matches(qa: Any, run: Any) -> tuple[int, int]:
    """Return (right_values, wrong_values) for schema-value comparison."""
    qa, run = _align_for_compare(qa, run)

    if _values_equivalent(qa, run):
        return 1, 0

    if type(qa) is not type(run):
        return 0, 1

    if isinstance(qa, dict):
        qa_map = {
            _normalize_field_key(k): k
            for k in qa.keys()
            if k not in IGNORED_META_KEYS
        }
        run_map = {
            _normalize_field_key(k): k
            for k in run.keys()
            if k not in IGNORED_META_KEYS
        }

        right = 0
        wrong = 0

        for missing_norm in set(qa_map.keys()) - set(run_map.keys()):
            _ = missing_norm
            wrong += 1

        for shared_norm in set(qa_map.keys()) & set(run_map.keys()):
            qa_key = qa_map[shared_norm]
            run_key = run_map[shared_norm]
            r, w = _count_value_matches(qa[qa_key], run[run_key])
            right += r
            wrong += w

        return right, wrong

    if isinstance(qa, list):
        right = 0
        wrong = 0

        if len(qa) != len(run):
            wrong += abs(len(qa) - len(run))

        for idx in range(min(len(qa), len(run))):
            r, w = _count_value_matches(qa[idx], run[idx])
            right += r
            wrong += w

        return right, wrong

    if qa == run:
        return 1, 0
    return 0, 1


def _normalize_field_key(key: str) -> str:
    # Align keys like `Agency` vs `agency` and `blkt_number` vs `BlktNumber`.
    return "".join(ch for ch in str(key).lower() if ch.isalnum())


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    if isinstance(value, (list, dict)):
        return len(value) == 0
    return False


def _normalize_scalar(value: Any) -> Any:
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return None

        # Date normalization: MM/DD/YYYY -> YYYY-MM-DD
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
        if m:
            mm, dd, yyyy = m.groups()
            return f"{yyyy}-{int(mm):02d}-{int(dd):02d}"

        m2 = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
        if m2:
            yyyy, mm, dd = m2.groups()
            return f"{int(yyyy):04d}-{int(mm):02d}-{int(dd):02d}"

        # Numeric normalization: "10,000" -> 10000 ; "2000" -> 2000
        numeric = s.replace(",", "")
        if re.fullmatch(r"[-+]?\d+", numeric):
            try:
                return int(numeric)
            except ValueError:
                return numeric

        if re.fullmatch(r"[-+]?\d+\.\d+", numeric):
            try:
                return float(numeric)
            except ValueError:
                return numeric

        return s

    if isinstance(value, (int, float, bool)):
        return value

    return value


def _values_equivalent(left: Any, right: Any) -> bool:
    # Treat null/blank/empty as equivalent when comparing schema values.
    if _is_blank(left) and _is_blank(right):
        return True

    # Only normalize scalar values here; containers are handled recursively.
    if isinstance(left, (dict, list)) or isinstance(right, (dict, list)):
        return False

    return _normalize_scalar(left) == _normalize_scalar(right)


def _build_header_map(headers: list[str]) -> dict[str, int]:
    return {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h is not None}


def _cell_value(values: tuple[Any, ...], header_map: dict[str, int], name: str) -> Any:
    idx = header_map.get(name.lower())
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compare QA and run JSON from Excel rows.")
    parser.add_argument("--excel", type=Path, required=True, help="Path to .xlsx file")
    parser.add_argument("--sheet", type=str, default=None, help="Sheet name (default: active sheet)")

    parser.add_argument("--qa-col", type=str, default="qa_json_path", help="Column name for QA JSON path")
    parser.add_argument("--run-col", type=str, default="run_json_path", help="Column name for run JSON path")
    parser.add_argument("--qa-json-col", type=str, default="qa_json", help="Column name for QA JSON text")
    parser.add_argument("--run-json-col", type=str, default="run_json", help="Column name for run JSON text")
    parser.add_argument("--key-col", type=str, default="file_name", help="Column name for key/filename")
    parser.add_argument(
        "--blob-uri-col",
        type=str,
        default="blob_uri",
        help="Column name for blob URI in Excel",
    )
    parser.add_argument(
        "--final-extraction-col",
        type=str,
        default="final_extraction",
        help="Column name for QA corrected JSON text in Excel",
    )

    parser.add_argument("--qa-dir", type=Path, default=None, help="Fallback QA JSON directory")
    parser.add_argument("--run-dir", type=Path, default=None, help="Fallback run JSON directory")

    parser.add_argument("--max-diffs", type=int, default=20, help="Max top differences to keep per row")
    parser.add_argument("--out-csv", type=Path, default=Path("data/outputs/acord140_diff_report.csv"))
    parser.add_argument("--out-json", type=Path, default=Path("data/outputs/acord140_diff_report.json"))
    return parser.parse_args()


def _resolve_row_sources(
    row_values: tuple[Any, ...],
    header_map: dict[str, int],
    qa_col: str,
    run_col: str,
    qa_json_col: str,
    run_json_col: str,
    key_col: str,
    blob_uri_col: str,
    final_extraction_col: str,
    qa_dir: Path | None,
    run_dir: Path | None,
) -> tuple[Any, Any, str, str, str]:
    qa_path_value = _cell_value(row_values, header_map, qa_col)
    run_path_value = _cell_value(row_values, header_map, run_col)
    qa_json_text = _cell_value(row_values, header_map, qa_json_col)
    run_json_text = _cell_value(row_values, header_map, run_json_col)
    key_value = _cell_value(row_values, header_map, key_col)
    blob_uri_value = _cell_value(row_values, header_map, blob_uri_col)
    final_extraction_value = _cell_value(row_values, header_map, final_extraction_col)

    # Prefer explicit blob_uri/final_extraction columns when provided.
    if final_extraction_value is not None and str(final_extraction_value).strip():
        qa_json_text = final_extraction_value
    if blob_uri_value is not None and str(blob_uri_value).strip():
        key_value = blob_uri_value

    key_text = str(key_value).strip() if key_value is not None else ""

    if qa_path_value and run_path_value:
        qa_source = str(qa_path_value).strip()
        run_source = str(run_path_value).strip()
        return _read_json_from_path(qa_source), _read_json_from_path(run_source), key_text, qa_source, run_source

    if qa_json_text and run_json_text:
        qa_source = f"cell:{qa_json_col}"
        run_source = f"cell:{run_json_col}"
        return _read_json_from_cell(str(qa_json_text)), _read_json_from_cell(str(run_json_text)), key_text, qa_source, run_source

    # New format: blob_uri + final_extraction JSON text + --run-dir
    if qa_json_text and key_text and run_dir:
        run_path = _resolve_run_json_from_blob_uri(run_dir, key_text)
        qa_source = f"cell:{final_extraction_col}"
        run_source = str(run_path)
        return _read_json_from_cell(str(qa_json_text)), _read_json_from_path(run_source), key_text, qa_source, run_source

    if key_text and qa_dir and run_dir:
        qa_path = _path_join_json(qa_dir, key_text)
        run_path = _path_join_json(run_dir, key_text)
        return _read_json_from_path(str(qa_path)), _read_json_from_path(str(run_path)), key_text, str(qa_path), str(run_path)

    raise ValueError(
        "Could not resolve row sources. Provide either path columns, JSON text columns, "
        "or key column with --qa-dir and --run-dir."
    )


def main() -> None:
    args = parse_args()

    if not args.excel.exists():
        raise SystemExit(f"Excel file not found: {args.excel}")

    wb = load_workbook(filename=args.excel, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        raise SystemExit("Excel sheet is empty.")

    if not headers:
        raise SystemExit("Header row is empty.")

    header_map = _build_header_map([str(h) if h is not None else "" for h in headers])

    results: list[RowResult] = []
    details: list[dict[str, Any]] = []

    total_rows = 0
    matched_rows = 0
    total_diffs = 0

    for excel_row_number, row_values in enumerate(rows_iter, start=2):
        if not row_values or all(v is None or str(v).strip() == "" for v in row_values):
            continue

        total_rows += 1

        try:
            qa_json, run_json, key_text, qa_source, run_source = _resolve_row_sources(
                row_values=row_values,
                header_map=header_map,
                qa_col=args.qa_col,
                run_col=args.run_col,
                qa_json_col=args.qa_json_col,
                run_json_col=args.run_json_col,
                key_col=args.key_col,
                blob_uri_col=args.blob_uri_col,
                final_extraction_col=args.final_extraction_col,
                qa_dir=args.qa_dir,
                run_dir=args.run_dir,
            )

            qa_json = _to_json(qa_json)
            run_json = _to_json(run_json)
            diffs = _collect_diffs(qa_json, run_json)
            diff_count = len(diffs)
            right_values, wrong_values = _count_value_matches(qa_json, run_json)
            total_diffs += diff_count
            if diff_count == 0:
                matched_rows += 1

            top = diffs[: max(0, args.max_diffs)]
            joined_top = " | ".join(top)

            key = key_text or _normalize_key(Path(qa_source).name)
            results.append(
                RowResult(
                    row_number=excel_row_number,
                    key=key,
                    qa_source=qa_source,
                    run_source=run_source,
                    diff_count=diff_count,
                    wrong_values=wrong_values,
                    right_values=right_values,
                    status="matched" if diff_count == 0 else "different",
                    top_differences=joined_top,
                )
            )

            details.append(
                {
                    "row_number": excel_row_number,
                    "key": key,
                    "qa_source": qa_source,
                    "run_source": run_source,
                    "diff_count": diff_count,
                    "wrong_values": wrong_values,
                    "right_values": right_values,
                    "status": "matched" if diff_count == 0 else "different",
                    "differences": diffs,
                }
            )
        except Exception as exc:
            results.append(
                RowResult(
                    row_number=excel_row_number,
                    key=str(_cell_value(row_values, header_map, args.key_col) or ""),
                    qa_source=str(_cell_value(row_values, header_map, args.qa_col) or ""),
                    run_source=str(_cell_value(row_values, header_map, args.run_col) or ""),
                    diff_count=-1,
                    wrong_values=0,
                    right_values=0,
                    status="error",
                    top_differences=str(exc),
                )
            )
            details.append(
                {
                    "row_number": excel_row_number,
                    "status": "error",
                    "error": str(exc),
                }
            )

    args.out_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.out_csv.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "row_number",
            "key",
            "qa_source",
            "run_source",
            "diff_count",
            "wrong_values",
            "right_values",
            "status",
            "top_differences",
        ])
        for row in results:
            writer.writerow([
                row.row_number,
                row.key,
                row.qa_source,
                row.run_source,
                row.diff_count,
                row.wrong_values,
                row.right_values,
                row.status,
                row.top_differences,
            ])

    args.out_json.parent.mkdir(parents=True, exist_ok=True)
    summary = {
        "excel": str(args.excel),
        "sheet": ws.title,
        "total_rows": total_rows,
        "matched_rows": matched_rows,
        "different_rows": sum(1 for r in results if r.status == "different"),
        "error_rows": sum(1 for r in results if r.status == "error"),
        "total_differences": total_diffs,
    }
    args.out_json.write_text(
        json.dumps({"summary": summary, "results": details}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("Comparison complete")
    print(f"Sheet: {ws.title}")
    print(f"Rows processed: {total_rows}")
    print(f"Matched rows: {matched_rows}")
    print(f"Different rows: {summary['different_rows']}")
    print(f"Error rows: {summary['error_rows']}")
    print(f"Total differences: {total_diffs}")
    print(f"CSV report: {args.out_csv}")
    print(f"JSON report: {args.out_json}")


if __name__ == "__main__":
    main()
